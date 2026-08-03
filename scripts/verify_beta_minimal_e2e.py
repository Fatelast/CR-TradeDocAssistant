"""执行 M5 Beta 合成样本的离线导入、审核模拟与导出闭环。"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKER_SOURCE = PROJECT_ROOT / "workers" / "excel-worker" / "src"
sys.path.insert(0, str(WORKER_SOURCE))

from protocol import PROTOCOL_VERSION, handle_line
from workbook import parse_workbook


class BetaAcceptanceError(Exception):
    """表示 Beta 演示样本技术闭环失败。"""


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _request(action: str, payload: dict[str, object], request_id: str) -> dict[str, Any]:
    response = handle_line(
        json.dumps(
            {
                "protocolVersion": PROTOCOL_VERSION,
                "id": request_id,
                "type": "request",
                "action": action,
                "payload": payload,
            },
        ),
    )
    if response.get("type") != "completed":
        error = response.get("error", {})
        raise BetaAcceptanceError(
            f"{action} 失败：{error.get('code', 'UNKNOWN')}"
        )
    return response["data"]


def _complete_task(task_detail: dict[str, Any], sample_id: str) -> dict[str, Any]:
    task_id = task_detail["task"]["taskId"]
    current = _request(
        "process_translation_batch",
        {"taskId": task_id, "batchSize": 200},
        f"{sample_id}-process",
    )
    for index, row in enumerate(current["rows"], start=1):
        if row["status"] in {"completed", "ignored"}:
            continue
        current = _request(
            "update_translation_row",
            {
                "taskId": task_id,
                "rowId": row["rowId"],
                "translation": f"Beta 审核译文 {index}",
            },
            f"{sample_id}-review-{index}",
        )
    if current["task"]["status"] != "completed":
        raise BetaAcceptanceError(f"{sample_id} 仍有未完成译文行")
    return current


def _assert_preservation(
    source_path: Path,
    output_path: Path,
    expected: dict[str, Any],
) -> None:
    """校验导出副本保留清单声明的公式与原有单元格样式。"""

    if not (
        expected["mustPreserveFormulas"]
        or expected["mustPreserveStyles"]
    ):
        return
    source_workbook = load_workbook(
        source_path,
        data_only=False,
        keep_links=False,
    )
    output_workbook = load_workbook(
        output_path,
        data_only=False,
        keep_links=False,
    )
    try:
        for source_sheet in source_workbook.worksheets:
            output_sheet = output_workbook[source_sheet.title]
            for row in source_sheet.iter_rows():
                for source_cell in row:
                    output_cell = output_sheet[source_cell.coordinate]
                    if (
                        expected["mustPreserveFormulas"]
                        and isinstance(source_cell.value, str)
                        and source_cell.value.startswith("=")
                        and output_cell.value != source_cell.value
                    ):
                        raise BetaAcceptanceError(
                            f"{source_path.name} 的公式 {source_sheet.title}!"
                            f"{source_cell.coordinate} 未被保留"
                        )
                    if (
                        expected["mustPreserveStyles"]
                        and source_cell.style_id != output_cell.style_id
                    ):
                        raise BetaAcceptanceError(
                            f"{source_path.name} 的样式 {source_sheet.title}!"
                            f"{source_cell.coordinate} 未被保留"
                        )
    finally:
        source_workbook.close()
        output_workbook.close()

def _verify_sample(
    sample: dict[str, Any],
    manifest_directory: Path,
    output_directory: Path,
) -> dict[str, Any]:
    output_directory.mkdir(parents=True, exist_ok=True)
    expected = sample["expected"]
    sample_path = manifest_directory / sample["file"]
    source_hash_before = _sha256(sample_path)
    detail = _request(
        "create_translation_task",
        {
            "filePath": str(sample_path),
            "sheetName": expected["sourceSheet"],
            "headerRow": expected["headerRow"],
            "sourceColumn": expected["sourceColumn"],
            "containerColumn": None,
            "targetColumn": expected["targetColumn"],
        },
        f"{sample['id']}-create",
    )
    completed = _complete_task(detail, sample["id"])
    task_id = completed["task"]["taskId"]
    preflight = _request(
        "preflight_export",
        {"taskId": task_id},
        f"{sample['id']}-preflight",
    )
    if not preflight["ready"]:
        raise BetaAcceptanceError(f"{sample['id']} 导出预检未就绪")
    output_path = output_directory / f"{sample['id']}-output.xlsx"
    exported = _request(
        "export_translation_task",
        {"taskId": task_id, "outputPath": str(output_path)},
        f"{sample['id']}-export",
    )
    if not exported["validated"] or not output_path.is_file():
        raise BetaAcceptanceError(f"{sample['id']} 未生成可验证的输出文件")
    if _sha256(sample_path) != source_hash_before:
        raise BetaAcceptanceError(f"{sample['id']} 的源文件在验收中被改变")
    _assert_preservation(sample_path, output_path, expected)
    output_data = parse_workbook(str(output_path))
    if not any(
        sheet["name"] == expected["sourceSheet"]
        for sheet in output_data["sheets"]
    ):
        raise BetaAcceptanceError(f"{sample['id']} 的输出缺少目标工作表")
    return {
        "id": sample["id"],
        "sourceSha256": source_hash_before,
        "outputFile": output_path.name,
        "writtenRows": exported["writtenRows"],
        "createdTargetColumn": exported["createdTargetColumn"],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--manifest",
        type=Path,
        default=PROJECT_ROOT / "resources" / "samples" / "m5-beta-minimal-manifest.json",
        help="Beta 合成样本清单路径",
    )
    arguments = parser.parse_args()
    manifest_path = arguments.manifest.resolve()
    previous_data_directory = os.environ.get("RUS_TRADE_DATA_DIR")
    try:
        document = json.loads(manifest_path.read_text(encoding="utf-8"))
        samples = document["samples"]
        if not isinstance(samples, list) or len(samples) != 3:
            raise BetaAcceptanceError("Beta 清单必须恰好包含 3 份样本")
        with tempfile.TemporaryDirectory(prefix="cr-trade-beta-") as directory:
            temporary_directory = Path(directory)
            os.environ["RUS_TRADE_DATA_DIR"] = str(temporary_directory / "data")
            results = [
                _verify_sample(
                    sample,
                    manifest_path.parent,
                    temporary_directory / "outputs",
                )
                for sample in samples
            ]
    except (BetaAcceptanceError, KeyError, OSError, ValueError) as error:
        print(json.dumps({"passed": False, "error": str(error)}, ensure_ascii=False))
        return 1
    finally:
        if previous_data_directory is None:
            os.environ.pop("RUS_TRADE_DATA_DIR", None)
        else:
            os.environ["RUS_TRADE_DATA_DIR"] = previous_data_directory

    print(json.dumps({"passed": True, "samples": results}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())