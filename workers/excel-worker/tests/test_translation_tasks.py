"""M2—M4 离线翻译任务、审核、本地数据与安全导出测试。"""

from __future__ import annotations

import json
import os
import shutil
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

WORKER_SRC = Path(__file__).resolve().parents[1] / "src"
SAMPLES = Path(__file__).resolve().parents[3] / "resources" / "samples"
sys.path.insert(0, str(WORKER_SRC))

from database import connect_database, resolve_database_path  # noqa: E402
from protocol import PROTOCOL_VERSION, handle_line  # noqa: E402
from translation import (  # noqa: E402
    ProtectionError,
    protect_text,
    restore_text,
)


class TranslationTaskTestCase(unittest.TestCase):
    def setUp(self) -> None:
        test_data_root = Path(__file__).resolve().parents[1] / ".test-data"
        test_data_root.mkdir(parents=True, exist_ok=True)
        self.temporary_directory = tempfile.TemporaryDirectory(
            dir=test_data_root,
        )
        self.test_directory = Path(self.temporary_directory.name)
        self.previous_data_directory = os.environ.get("RUS_TRADE_DATA_DIR")
        os.environ["RUS_TRADE_DATA_DIR"] = self.temporary_directory.name

    def tearDown(self) -> None:
        if self.previous_data_directory is None:
            os.environ.pop("RUS_TRADE_DATA_DIR", None)
        else:
            os.environ["RUS_TRADE_DATA_DIR"] = self.previous_data_directory
        self.temporary_directory.cleanup()

    def request(
        self,
        action: str,
        payload: dict[str, object],
    ) -> dict[str, object]:
        return handle_line(
            json.dumps(
                {
                    "protocolVersion": PROTOCOL_VERSION,
                    "id": f"req_{action}",
                    "type": "request",
                    "action": action,
                    "payload": payload,
                }
            )
        )

    def create_standard_task(
        self,
        file_path: Path | None = None,
        target_column: int | None = 4,
    ) -> dict[str, object]:
        response = self.request(
            "create_translation_task",
            {
                "filePath": str(file_path or SAMPLES / "m1-standard.xlsx"),
                "sheetName": "问题反馈",
                "headerRow": 2,
                "sourceColumn": 3,
                "containerColumn": 1,
                "targetColumn": target_column,
            },
        )
        self.assertEqual(response["type"], "completed")
        return response["data"]

    def complete_task(self, detail: dict[str, object]) -> dict[str, object]:
        task_id = detail["task"]["taskId"]
        processed = self.request(
            "process_translation_batch",
            {"taskId": task_id, "batchSize": 50},
        )
        self.assertEqual(processed["type"], "completed")
        current = processed["data"]
        for index, row in enumerate(current["rows"]):
            if row["status"] in {"completed", "ignored"}:
                continue
            updated = self.request(
                "update_translation_row",
                {
                    "taskId": task_id,
                    "rowId": row["rowId"],
                    "translation": f"审核译文 {index + 1}",
                },
            )
            self.assertEqual(updated["type"], "completed")
            current = updated["data"]
        self.assertEqual(current["task"]["status"], "completed")
        return current

    def test_database_migration_and_builtin_glossary(self) -> None:
        with connect_database() as connection:
            version = connection.execute("PRAGMA user_version").fetchone()[0]
            term_count = connection.execute(
                "SELECT COUNT(*) FROM glossary_terms"
            ).fetchone()[0]
            task_columns = {
                row["name"]
                for row in connection.execute("PRAGMA table_info(tasks)")
            }

        self.assertEqual(version, 3)
        self.assertEqual(term_count, 5)
        self.assertIn("source_file_sha256", task_columns)
        self.assertIn("rerun_of_task_id", task_columns)

    def test_schema_v2_data_survives_v3_migration(self) -> None:
        detail = self.create_standard_task()
        task_id = detail["task"]["taskId"]
        with connect_database() as connection:
            term_count = connection.execute(
                "SELECT COUNT(*) FROM glossary_terms"
            ).fetchone()[0]

        connection = sqlite3.connect(resolve_database_path())
        try:
            connection.executescript(
                """
                DROP INDEX IF EXISTS idx_glossary_term_key;
                DROP INDEX IF EXISTS idx_glossary_category;
                DROP INDEX IF EXISTS idx_tasks_history;
                DROP INDEX IF EXISTS idx_cache_expiry;
                DROP INDEX IF EXISTS idx_task_exports_task;
                DROP INDEX IF EXISTS idx_task_exports_status;
                DROP TABLE IF EXISTS task_exports;
                DROP TABLE IF EXISTS app_settings;
                ALTER TABLE glossary_terms DROP COLUMN term_key;
                ALTER TABLE tasks DROP COLUMN rerun_of_task_id;
                PRAGMA user_version = 2;
                """
            )
            connection.commit()
        finally:
            connection.close()

        with connect_database() as connection:
            self.assertEqual(
                connection.execute("PRAGMA user_version").fetchone()[0],
                3,
            )
            self.assertIsNotNone(
                connection.execute(
                    "SELECT id FROM tasks WHERE id = ?",
                    (task_id,),
                ).fetchone()
            )
            self.assertEqual(
                connection.execute(
                    "SELECT COUNT(*) FROM glossary_terms"
                ).fetchone()[0],
                term_count,
            )
            self.assertEqual(
                connection.execute(
                    "SELECT COUNT(*) FROM glossary_terms "
                    "WHERE term_key IS NULL"
                ).fetchone()[0],
                0,
            )
    def test_task_batch_manual_update_cache_and_recovery(self) -> None:
        detail = self.create_standard_task()
        task = detail["task"]

        self.assertEqual(task["status"], "draft")
        self.assertEqual(task["totalRows"], 5)
        self.assertEqual(task["pendingRows"], 4)
        self.assertEqual(task["completedRows"], 1)
        self.assertTrue(task["sourceFingerprintAvailable"])

        processed = self.request(
            "process_translation_batch",
            {"taskId": task["taskId"], "batchSize": 50},
        )
        self.assertEqual(processed["type"], "completed")
        processed_detail = processed["data"]
        self.assertEqual(processed_detail["task"]["status"], "awaiting_manual")
        self.assertEqual(processed_detail["task"]["candidateRows"], 4)
        self.assertEqual(processed_detail["task"]["completedRows"], 1)

        open_rows = [
            row
            for row in processed_detail["rows"]
            if row["status"] != "completed"
        ]
        first_row = open_rows[0]
        self.assertEqual(
            first_row["initialTranslation"],
            first_row["translation"],
        )
        for index, row in enumerate(open_rows):
            updated = self.request(
                "update_translation_row",
                {
                    "taskId": task["taskId"],
                    "rowId": row["rowId"],
                    "translation": f"人工确认译文 {index + 1}",
                    "saveToCache": index == 0,
                },
            )
            self.assertEqual(updated["type"], "completed")

        final_detail = updated["data"]
        self.assertEqual(final_detail["task"]["status"], "completed")
        self.assertEqual(final_detail["task"]["completedRows"], 5)

        second_detail = self.create_standard_task()
        second_task_id = second_detail["task"]["taskId"]
        second_processed = self.request(
            "process_translation_batch",
            {"taskId": second_task_id},
        )
        cached_row = next(
            row
            for row in second_processed["data"]["rows"]
            if row["sourceText"] == first_row["sourceText"]
        )
        self.assertEqual(cached_row["candidateSource"], "cache")

        incomplete = self.request(
            "list_translation_tasks",
            {"includeCompleted": False},
        )
        self.assertEqual(incomplete["type"], "completed")
        self.assertEqual(len(incomplete["data"]["tasks"]), 1)
        self.assertEqual(
            incomplete["data"]["tasks"][0]["taskId"],
            second_task_id,
        )

    def test_exact_glossary_and_task_state_changes(self) -> None:
        term = self.request(
            "upsert_glossary_term",
            {
                "sourceText": "Повреждена внешняя упаковка",
                "targetText": "外包装受损",
                "exactMatch": True,
            },
        )
        self.assertEqual(term["type"], "completed")

        detail = self.create_standard_task()
        task_id = detail["task"]["taskId"]
        paused = self.request(
            "change_translation_task_state",
            {"taskId": task_id, "action": "pause"},
        )
        self.assertEqual(paused["data"]["task"]["status"], "paused")

        resumed = self.request(
            "change_translation_task_state",
            {"taskId": task_id, "action": "resume"},
        )
        self.assertEqual(resumed["data"]["task"]["status"], "running")

        processed = self.request(
            "process_translation_batch",
            {"taskId": task_id},
        )
        exact_row = next(
            row
            for row in processed["data"]["rows"]
            if row["sourceText"] == "Повреждена внешняя упаковка"
        )
        self.assertEqual(exact_row["translation"], "外包装受损")
        self.assertEqual(exact_row["candidateSource"], "glossary")

        cancelled = self.request(
            "change_translation_task_state",
            {"taskId": task_id, "action": "cancel"},
        )
        self.assertEqual(cancelled["data"]["task"]["status"], "cancelled")

    def test_review_restore_ignore_and_rematch(self) -> None:
        detail = self.create_standard_task()
        task_id = detail["task"]["taskId"]
        processed = self.request(
            "process_translation_batch",
            {"taskId": task_id},
        )["data"]
        row = next(
            item for item in processed["rows"] if item["status"] == "candidate"
        )
        initial_translation = row["initialTranslation"]

        self.request(
            "update_translation_row",
            {
                "taskId": task_id,
                "rowId": row["rowId"],
                "translation": "人工修改",
            },
        )
        restored = self.request(
            "review_translation_row",
            {
                "taskId": task_id,
                "rowId": row["rowId"],
                "action": "restore_initial",
            },
        )
        restored_row = next(
            item
            for item in restored["data"]["rows"]
            if item["rowId"] == row["rowId"]
        )
        self.assertEqual(restored_row["translation"], initial_translation)
        self.assertEqual(restored_row["status"], "candidate")
        self.assertFalse(restored_row["userModified"])

        ignored = self.request(
            "review_translation_row",
            {
                "taskId": task_id,
                "rowId": row["rowId"],
                "action": "ignore",
            },
        )
        ignored_row = next(
            item
            for item in ignored["data"]["rows"]
            if item["rowId"] == row["rowId"]
        )
        self.assertEqual(ignored_row["status"], "ignored")

        rematched = self.request(
            "review_translation_row",
            {
                "taskId": task_id,
                "rowId": row["rowId"],
                "action": "rematch",
            },
        )
        rematched_row = next(
            item
            for item in rematched["data"]["rows"]
            if item["rowId"] == row["rowId"]
        )
        self.assertEqual(rematched_row["status"], "candidate")
        self.assertFalse(rematched_row["userModified"])

    def test_export_existing_target_and_validate(self) -> None:
        completed = self.complete_task(self.create_standard_task())
        task_id = completed["task"]["taskId"]
        preflight = self.request("preflight_export", {"taskId": task_id})

        self.assertEqual(preflight["type"], "completed")
        self.assertTrue(preflight["data"]["ready"])
        self.assertFalse(preflight["data"]["createsTargetColumn"])
        self.assertEqual(preflight["data"]["targetColumnLetter"], "D")
        self.assertEqual(preflight["data"]["writableRows"], 4)

        output_path = self.test_directory / "translated-existing.xlsx"
        exported = self.request(
            "export_translation_task",
            {"taskId": task_id, "outputPath": str(output_path)},
        )
        self.assertEqual(exported["type"], "completed")
        self.assertTrue(exported["data"]["validated"])
        self.assertEqual(exported["data"]["writtenRows"], 4)

        workbook = load_workbook(output_path, data_only=False)
        try:
            worksheet = workbook["问题反馈"]
            self.assertEqual(worksheet["D3"].value, "审核译文 1")
            self.assertEqual(worksheet["D4"].value, "纸箱缺少标签")
            self.assertEqual(worksheet["E3"].value, "=LEN(C3)")
        finally:
            workbook.close()

    def test_export_appends_target_column_and_copies_style(self) -> None:
        completed = self.complete_task(
            self.create_standard_task(target_column=None)
        )
        task_id = completed["task"]["taskId"]
        preflight = self.request("preflight_export", {"taskId": task_id})
        self.assertEqual(preflight["data"]["targetColumnLetter"], "F")
        self.assertTrue(preflight["data"]["createsTargetColumn"])

        output_path = self.test_directory / "translated-new-column.xlsx"
        exported = self.request(
            "export_translation_task",
            {"taskId": task_id, "outputPath": str(output_path)},
        )
        self.assertEqual(exported["type"], "completed")
        self.assertTrue(exported["data"]["createdTargetColumn"])

        workbook = load_workbook(output_path, data_only=False)
        try:
            worksheet = workbook["问题反馈"]
            self.assertEqual(worksheet["F2"].value, "中文翻译")
            self.assertEqual(worksheet["F3"].value, "审核译文 1")
            self.assertEqual(worksheet["F3"].font.name, worksheet["C3"].font.name)
            self.assertEqual(worksheet["E3"].value, "=LEN(C3)")
        finally:
            workbook.close()

    def test_export_blocks_incomplete_and_changed_source(self) -> None:
        source_copy = self.test_directory / "source-copy.xlsx"
        shutil.copy2(SAMPLES / "m1-standard.xlsx", source_copy)
        detail = self.create_standard_task(file_path=source_copy)
        task_id = detail["task"]["taskId"]

        incomplete = self.request("preflight_export", {"taskId": task_id})
        self.assertEqual(incomplete["type"], "error")
        self.assertEqual(
            incomplete["error"]["code"],
            "EXPORT_TASK_INCOMPLETE",
        )

        completed = self.complete_task(detail)
        workbook = load_workbook(source_copy)
        try:
            workbook["问题反馈"]["B3"] = "changed"
            workbook.save(source_copy)
        finally:
            workbook.close()

        changed = self.request(
            "preflight_export",
            {"taskId": completed["task"]["taskId"]},
        )
        self.assertEqual(changed["type"], "error")
        self.assertEqual(changed["error"]["code"], "SOURCE_FILE_CHANGED")

    def test_export_blocks_restricted_workbook(self) -> None:
        created = self.request(
            "create_translation_task",
            {
                "filePath": str(SAMPLES / "m1-risky.xlsx"),
                "sheetName": "风险对象",
                "headerRow": 1,
                "sourceColumn": 3,
                "containerColumn": 1,
                "targetColumn": 4,
            },
        )
        self.assertEqual(created["type"], "completed")
        self.assertTrue(created["data"]["task"]["sourceRestricted"])
        completed = self.complete_task(created["data"])

        preflight = self.request(
            "preflight_export",
            {"taskId": completed["task"]["taskId"]},
        )
        self.assertEqual(preflight["type"], "error")
        self.assertEqual(preflight["error"]["code"], "EXPORT_RESTRICTED")

    def test_m4_settings_and_glossary_xlsx_round_trip(self) -> None:
        settings = self.request("get_app_settings", {})
        self.assertEqual(settings["type"], "completed")
        self.assertEqual(settings["data"]["batchSize"], 50)
        self.assertEqual(settings["data"]["retention"]["taskDays"], 90)

        updated = self.request(
            "update_app_settings",
            {
                "defaultOutputDirectory": str(self.test_directory),
                "batchSize": 25,
                "logLevel": "error",
            },
        )
        self.assertEqual(updated["data"]["batchSize"], 25)
        settings_path = self.test_directory / "settings.json"
        exported_settings = self.request(
            "export_app_settings",
            {"outputPath": str(settings_path)},
        )
        self.assertEqual(exported_settings["type"], "completed")
        with settings_path.open(encoding="utf-8") as source:
            settings_document = json.load(source)
        self.assertEqual(settings_document["format"], "cr-trade-settings")

        glossary_path = self.test_directory / "glossary.xlsx"
        exported_glossary = self.request(
            "export_glossary",
            {"outputPath": str(glossary_path)},
        )
        self.assertEqual(exported_glossary["type"], "completed")
        workbook = load_workbook(glossary_path)
        try:
            worksheet = workbook.active
            worksheet.cell(2, 2, "已更新译文")
            worksheet.append(
                ["Новый термин", "新术语", "测试", "是", "否", "是", ""]
            )
            workbook.save(glossary_path)
        finally:
            workbook.close()

        preflight = self.request(
            "preflight_glossary_import",
            {"filePath": str(glossary_path)},
        )
        self.assertEqual(preflight["type"], "completed")
        self.assertEqual(preflight["data"]["createdRows"], 1)
        self.assertEqual(preflight["data"]["updatedRows"], 1)
        applied = self.request(
            "apply_glossary_import",
            {
                "filePath": str(glossary_path),
                "expectedSha256": preflight["data"]["fileSha256"],
            },
        )
        self.assertEqual(applied["type"], "completed")
        listed = self.request(
            "list_glossary_terms",
            {"search": "Новый термин"},
        )
        self.assertEqual(listed["data"]["total"], 1)
        formula_term = self.request(
            "upsert_glossary_term",
            {"sourceText": "=HYPERLINK()", "targetText": "危险"},
        )
        self.assertEqual(formula_term["error"]["code"], "INVALID_MESSAGE")

    def test_m4_history_rerun_export_and_confirmed_cleanup(self) -> None:
        detail = self.create_standard_task()
        completed = self.complete_task(detail)
        task_id = completed["task"]["taskId"]
        output_path = self.test_directory / "history-export.xlsx"
        exported = self.request(
            "export_translation_task",
            {"taskId": task_id, "outputPath": str(output_path)},
        )
        self.assertEqual(exported["type"], "completed")
        self.assertTrue(exported["data"]["exportId"])

        history = self.request(
            "list_task_history",
            {"search": "m1-standard", "page": 1, "pageSize": 10},
        )
        self.assertEqual(history["data"]["total"], 1)
        self.assertEqual(history["data"]["items"][0]["exportCount"], 1)
        history_detail = self.request(
            "get_task_history_detail",
            {"taskId": task_id},
        )
        self.assertEqual(
            history_detail["data"]["exports"][0]["status"],
            "completed",
        )

        rerun = self.request("create_rerun_task", {"taskId": task_id})
        self.assertEqual(rerun["type"], "completed")
        self.assertNotEqual(rerun["data"]["task"]["taskId"], task_id)
        self.assertEqual(rerun["data"]["task"]["rerunOfTaskId"], task_id)

        cleanup_request = {
            "scopes": ["selected_tasks"],
            "taskIds": [task_id],
            "cacheKeys": [],
        }
        preview = self.request("preview_data_cleanup", cleanup_request)
        self.assertEqual(preview["data"]["taskCount"], 1)
        self.assertEqual(preview["data"]["exportCount"], 1)
        with connect_database() as connection:
            connection.execute(
                "UPDATE tasks SET updated_at = ? WHERE id = ?",
                ("2099-01-01T00:00:00+00:00", task_id),
            )
        stale_cleanup = self.request(
            "run_data_cleanup",
            {**cleanup_request, "planHash": preview["data"]["planHash"]},
        )
        self.assertEqual(
            stale_cleanup["error"]["code"],
            "CLEANUP_PLAN_CHANGED",
        )
        preview = self.request("preview_data_cleanup", cleanup_request)
        cleanup = self.request(
            "run_data_cleanup",
            {**cleanup_request, "planHash": preview["data"]["planHash"]},
        )
        self.assertEqual(cleanup["type"], "completed")
        self.assertTrue(output_path.is_file())
        self.assertTrue((SAMPLES / "m1-standard.xlsx").is_file())
        deleted = self.request(
            "get_task_history_detail",
            {"taskId": task_id},
        )
        self.assertEqual(deleted["error"]["code"], "HISTORY_TASK_NOT_FOUND")
    def test_protected_fragments_round_trip_and_validation(self) -> None:
        source = (
            "Контейнер MSCU1234567, заказ PO-2026-001, "
            "дата 2026-07-30, сумма USD 1250.50"
        )
        protected, values = protect_text(source)

        self.assertNotIn("MSCU1234567", protected)
        self.assertEqual(restore_text(protected, values), source)

        with self.assertRaises(ProtectionError):
            restore_text(protected.replace("⟦P0001⟧", ""), values)


if __name__ == "__main__":
    unittest.main()
