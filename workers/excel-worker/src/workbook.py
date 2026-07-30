"""Excel 工作簿只读解析与 M1 导入预检。"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile, is_zipfile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5_000
MAX_COLUMNS = 200
HEADER_SCAN_ROWS = 12
PREVIEW_ROWS = 12
MAX_ARCHIVE_SIZE_BYTES = 200 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 5_000

LIMITS = {
    "maxFileSizeBytes": MAX_FILE_SIZE_BYTES,
    "maxRows": MAX_ROWS,
    "maxColumns": MAX_COLUMNS,
    "headerScanRows": HEADER_SCAN_ROWS,
    "previewRows": PREVIEW_ROWS,
}

RISK_MESSAGES = {
    "EXTERNAL_LINKS": "检测到外部链接，当前文件仅允许只读预览",
    "CHARTS": "检测到图表，当前文件仅允许只读预览",
    "DRAWINGS": "检测到绘图或图片对象，当前文件仅允许只读预览",
    "DATA_VALIDATION": "检测到数据验证，当前文件仅允许只读预览",
    "CONDITIONAL_FORMATTING": "检测到条件格式，当前文件仅允许只读预览",
    "MERGED_CELLS": "检测到合并单元格，请确认表头和目标列",
    "FORMULAS": "检测到公式；M1 仅展示公式，不会计算或修改",
    "HIDDEN_SHEETS": "工作簿包含隐藏工作表",
}

RESTRICTED_RISKS = {
    "EXTERNAL_LINKS",
    "CHARTS",
    "DRAWINGS",
    "DATA_VALIDATION",
    "CONDITIONAL_FORMATTING",
}

CONTAINER_KEYWORDS = (
    "箱号",
    "柜号",
    "集装箱",
    "container",
    "контейнер",
)
SOURCE_KEYWORDS = (
    "问题",
    "反馈",
    "异常",
    "备注",
    "描述",
    "issue",
    "feedback",
    "comment",
    "problem",
    "проблем",
    "замечан",
    "описан",
)
TARGET_KEYWORDS = (
    "中文",
    "翻译",
    "译文",
    "translation",
    "chinese",
    "перевод",
)


@dataclass(frozen=True)
class WorkbookError(Exception):
    """可映射到稳定 Worker 错误码的工作簿异常。"""

    code: str
    message: str


def _validate_file(file_path: Any) -> tuple[Path, int]:
    if not isinstance(file_path, str) or not file_path.strip():
        raise WorkbookError("INVALID_MESSAGE", "文件路径不能为空")

    path = Path(file_path).expanduser()
    if not path.is_file():
        raise WorkbookError("FILE_NOT_FOUND", "文件不存在")
    if path.suffix.lower() != ".xlsx":
        raise WorkbookError("FILE_UNSUPPORTED", "仅支持 .xlsx 文件")

    try:
        file_size = path.stat().st_size
    except PermissionError as error:
        raise WorkbookError("FILE_LOCKED", "文件被占用或无权读取") from error

    if file_size > MAX_FILE_SIZE_BYTES:
        raise WorkbookError(
            "FILE_TOO_LARGE",
            "文件超过 10 MiB 上限",
        )

    if not is_zipfile(path):
        try:
            magic = path.read_bytes()[:8]
        except PermissionError as error:
            raise WorkbookError("FILE_LOCKED", "文件被占用或无权读取") from error

        if magic == bytes.fromhex("D0CF11E0A1B11AE1"):
            raise WorkbookError("WORKBOOK_ENCRYPTED", "文件可能已加密")
        raise WorkbookError("WORKBOOK_OPEN_FAILED", "文件不是有效的 .xlsx 工作簿")

    return path.resolve(), file_size


def _add_risk(risks: dict[str, dict[str, str]], code: str) -> None:
    risks[code] = {
        "code": code,
        "severity": "restricted" if code in RESTRICTED_RISKS else "notice",
        "message": RISK_MESSAGES[code],
    }


def _scan_archive(path: Path) -> dict[str, dict[str, str]]:
    risks: dict[str, dict[str, str]] = {}

    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise WorkbookError(
                    "WORKBOOK_ARCHIVE_TOO_LARGE",
                    "工作簿内部文件数量超过安全上限",
                )

            unpacked_size = sum(entry.file_size for entry in entries)
            if unpacked_size > MAX_ARCHIVE_SIZE_BYTES:
                raise WorkbookError(
                    "WORKBOOK_ARCHIVE_TOO_LARGE",
                    "工作簿解压后超过 200 MiB 安全上限",
                )

            names = [entry.filename.lower() for entry in entries]
            if any(name.startswith("xl/externallinks/") for name in names):
                _add_risk(risks, "EXTERNAL_LINKS")
            if any(name.startswith("xl/charts/") for name in names):
                _add_risk(risks, "CHARTS")
            if any(name.startswith("xl/drawings/") for name in names):
                _add_risk(risks, "DRAWINGS")

            worksheet_names = [
                entry.filename
                for entry in entries
                if entry.filename.lower().startswith("xl/worksheets/sheet")
                and entry.filename.lower().endswith(".xml")
            ]
            patterns = {
                "DATA_VALIDATION": re.compile(
                    rb"<(?:[A-Za-z0-9_]+:)?dataValidations\b"
                ),
                "CONDITIONAL_FORMATTING": re.compile(
                    rb"<(?:[A-Za-z0-9_]+:)?conditionalFormatting\b"
                ),
                "MERGED_CELLS": re.compile(
                    rb"<(?:[A-Za-z0-9_]+:)?mergeCells\b"
                ),
                "FORMULAS": re.compile(rb"<(?:[A-Za-z0-9_]+:)?f(?:\s|>)"),
            }
            for worksheet_name in worksheet_names:
                xml = archive.read(worksheet_name)
                for code, pattern in patterns.items():
                    if code not in risks and pattern.search(xml):
                        _add_risk(risks, code)
    except BadZipFile as error:
        raise WorkbookError(
            "WORKBOOK_OPEN_FAILED",
            "无法读取工作簿压缩结构",
        ) from error

    return risks


def _open_workbook(path: Path):
    try:
        return load_workbook(
            filename=path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except PermissionError as error:
        raise WorkbookError("FILE_LOCKED", "文件被占用或无权读取") from error
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as error:
        raise WorkbookError("WORKBOOK_OPEN_FAILED", "无法读取 Excel 文件") from error


def _validate_sheet_size(worksheet) -> tuple[int, int]:
    """按非空值计算有效区域，避免仅有样式的空单元格触发误判。"""

    row_count = 0
    column_count = 0
    for row_number, row in enumerate(
        worksheet.iter_rows(values_only=True),
        start=1,
    ):
        populated_columns = [
            index
            for index, value in enumerate(row, start=1)
            if value is not None and (not isinstance(value, str) or value.strip())
        ]
        if not populated_columns:
            continue

        row_count = row_number
        column_count = max(column_count, max(populated_columns))
        if row_count > MAX_ROWS:
            raise WorkbookError(
                "ROW_LIMIT_EXCEEDED",
                f"工作表“{worksheet.title}”超过 5,000 行上限",
            )
        if column_count > MAX_COLUMNS:
            raise WorkbookError(
                "COLUMN_LIMIT_EXCEEDED",
                f"工作表“{worksheet.title}”超过 200 列上限",
            )

    return row_count, column_count


def _normalize_label(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value)).lower()


def _keyword_hits(labels: list[str]) -> int:
    keywords = CONTAINER_KEYWORDS + SOURCE_KEYWORDS + TARGET_KEYWORDS
    return sum(
        1
        for label in labels
        if any(keyword in label for keyword in keywords)
    )


def _header_candidates(worksheet, row_count: int, column_count: int) -> list[int]:
    if row_count <= 0 or column_count <= 0:
        return [1]

    candidates: list[tuple[int, int]] = []
    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=min(row_count, HEADER_SCAN_ROWS),
            min_col=1,
            max_col=column_count,
            values_only=True,
        ),
        start=1,
    ):
        labels = [_normalize_label(value) for value in row]
        populated = sum(1 for label in labels if label)
        if populated == 0:
            continue
        score = populated + (_keyword_hits(labels) * 4)
        candidates.append((row_number, score))

    if not candidates:
        return [1]

    ranked = sorted(candidates, key=lambda item: (-item[1], item[0]))
    recommended = ranked[0][0]
    alternatives = [row for row, _score in ranked[1:3]]
    return [recommended, *sorted(alternatives)]


def _serialize_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _get_worksheet(workbook, sheet_name: Any):
    if not isinstance(sheet_name, str) or sheet_name not in workbook.sheetnames:
        raise WorkbookError("SHEET_NOT_FOUND", "工作表不存在")
    return workbook[sheet_name]


def _validate_header_row(header_row: Any, row_count: int) -> int:
    if not isinstance(header_row, int) or isinstance(header_row, bool):
        raise WorkbookError("INVALID_MESSAGE", "表头行必须是整数")
    if header_row < 1 or header_row > max(row_count, 1):
        raise WorkbookError("HEADER_NOT_FOUND", "表头行超出工作表范围")
    return header_row


def _recommend_column(labels: list[str], keywords: tuple[str, ...]) -> int | None:
    for index, label in enumerate(labels, start=1):
        if any(keyword in label for keyword in keywords):
            return index
    return None


def parse_workbook(file_path: Any) -> dict[str, Any]:
    """执行文件限制、工作表结构和高风险特征预检。"""

    path, file_size = _validate_file(file_path)
    risks = _scan_archive(path)
    workbook = _open_workbook(path)

    try:
        sheets: list[dict[str, Any]] = []
        for worksheet in workbook.worksheets:
            row_count, column_count = _validate_sheet_size(worksheet)
            candidates = _header_candidates(
                worksheet,
                row_count,
                column_count,
            )
            sheets.append(
                {
                    "name": worksheet.title,
                    "state": worksheet.sheet_state,
                    "rowCount": row_count,
                    "columnCount": column_count,
                    "headerCandidateRows": candidates,
                    "recommendedHeaderRow": candidates[0],
                }
            )

        if not sheets:
            raise WorkbookError("WORKBOOK_OPEN_FAILED", "工作簿不包含工作表")
        if any(sheet["state"] != "visible" for sheet in sheets):
            _add_risk(risks, "HIDDEN_SHEETS")

        default_sheet = next(
            (sheet for sheet in sheets if sheet["state"] == "visible"),
            sheets[0],
        )
        risk_list = list(risks.values())
        return {
            "fileName": path.name,
            "fileSizeBytes": file_size,
            "defaultSheetName": default_sheet["name"],
            "sheets": sheets,
            "risks": risk_list,
            "restricted": any(
                risk["severity"] == "restricted" for risk in risk_list
            ),
            "limits": LIMITS,
        }
    finally:
        workbook.close()


def preview_sheet(
    file_path: Any,
    sheet_name: Any,
    header_row: Any,
) -> dict[str, Any]:
    """返回指定工作表的表头、列推荐和有限数据预览。"""

    path, _file_size = _validate_file(file_path)
    workbook = _open_workbook(path)

    try:
        worksheet = _get_worksheet(workbook, sheet_name)
        row_count, column_count = _validate_sheet_size(worksheet)
        header = _validate_header_row(header_row, row_count)
        if column_count <= 0:
            raise WorkbookError("HEADER_NOT_FOUND", "工作表没有可用列")

        header_values = next(
            worksheet.iter_rows(
                min_row=header,
                max_row=header,
                min_col=1,
                max_col=column_count,
                values_only=True,
            )
        )
        labels = [_normalize_label(value) for value in header_values]
        if not any(labels):
            raise WorkbookError("HEADER_NOT_FOUND", "所选行不包含表头")

        columns = [
            {
                "index": index,
                "letter": get_column_letter(index),
                "label": str(value).strip()
                if value is not None and str(value).strip()
                else f"未命名列 {get_column_letter(index)}",
            }
            for index, value in enumerate(header_values, start=1)
        ]

        rows: list[dict[str, Any]] = []
        preview_end = min(row_count, header + PREVIEW_ROWS)
        if preview_end > header:
            for excel_row, values in enumerate(
                worksheet.iter_rows(
                    min_row=header + 1,
                    max_row=preview_end,
                    min_col=1,
                    max_col=column_count,
                    values_only=True,
                ),
                start=header + 1,
            ):
                rows.append(
                    {
                        "excelRowNumber": excel_row,
                        "values": [_serialize_value(value) for value in values],
                    }
                )

        return {
            "sheetName": worksheet.title,
            "headerRow": header,
            "totalRows": max(row_count - header, 0),
            "columns": columns,
            "recommendations": {
                "containerColumn": _recommend_column(
                    labels,
                    CONTAINER_KEYWORDS,
                ),
                "sourceColumn": _recommend_column(labels, SOURCE_KEYWORDS),
                "targetColumn": _recommend_column(labels, TARGET_KEYWORDS),
            },
            "rows": rows,
        }
    finally:
        workbook.close()


def _optional_column(value: Any, error_code: str, message: str) -> int | None:
    if value is None:
        return None
    if not isinstance(value, int) or isinstance(value, bool):
        raise WorkbookError(error_code, message)
    if value < 1 or value > MAX_COLUMNS:
        raise WorkbookError(error_code, message)
    return value


def build_import_rows(
    file_path: Any,
    sheet_name: Any,
    header_row: Any,
    source_column: Any,
    container_column: Any = None,
    target_column: Any = None,
) -> dict[str, Any]:
    """根据用户确认的列生成稳定、可追踪的 M1 数据行。"""

    path, _file_size = _validate_file(file_path)
    source = _optional_column(
        source_column,
        "SOURCE_COLUMN_INVALID",
        "原文列无效",
    )
    if source is None:
        raise WorkbookError("SOURCE_COLUMN_INVALID", "必须选择原文列")
    container = _optional_column(
        container_column,
        "SOURCE_COLUMN_INVALID",
        "箱号列无效",
    )
    target = _optional_column(
        target_column,
        "TARGET_COLUMN_INVALID",
        "译文列无效",
    )

    workbook = _open_workbook(path)
    try:
        worksheet = _get_worksheet(workbook, sheet_name)
        row_count, column_count = _validate_sheet_size(worksheet)
        header = _validate_header_row(header_row, row_count)
        for column, code, message in (
            (source, "SOURCE_COLUMN_INVALID", "原文列超出工作表范围"),
            (container, "SOURCE_COLUMN_INVALID", "箱号列超出工作表范围"),
            (target, "TARGET_COLUMN_INVALID", "译文列超出工作表范围"),
        ):
            if column is not None and column > column_count:
                raise WorkbookError(code, message)

        requested_columns = [source]
        if container is not None:
            requested_columns.append(container)
        if target is not None:
            requested_columns.append(target)
        max_requested_column = max(requested_columns)

        rows: list[dict[str, Any]] = []
        for cells in worksheet.iter_rows(
            min_row=header + 1,
            max_row=row_count,
            min_col=1,
            max_col=max_requested_column,
            values_only=False,
        ):
            source_cell = cells[source - 1]
            source_value = source_cell.value
            if source_value is None or not str(source_value).strip():
                continue

            container_cell = cells[container - 1] if container else None
            target_cell = cells[target - 1] if target else None
            rows.append(
                {
                    "rowId": (
                        f"{worksheet.title}:{source_cell.row}:"
                        f"{get_column_letter(source)}"
                    ),
                    "excelRowNumber": source_cell.row,
                    "sourceCell": source_cell.coordinate,
                    "targetCell": target_cell.coordinate if target_cell else None,
                    "containerCell": (
                        container_cell.coordinate if container_cell else None
                    ),
                    "containerValue": _serialize_value(
                        container_cell.value if container_cell else None
                    ),
                    "sourceText": str(source_value),
                    "existingTarget": _serialize_value(
                        target_cell.value if target_cell else None
                    ),
                    "status": "pending",
                }
            )

        return {
            "sheetName": worksheet.title,
            "totalRows": len(rows),
            "rows": rows,
        }
    finally:
        workbook.close()
