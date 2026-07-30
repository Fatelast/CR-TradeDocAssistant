"""M3 翻译任务导出预检与 Excel 安全写回。"""

from __future__ import annotations

import hashlib
import json
import os
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

from database import connect_database
from workbook import MAX_COLUMNS, parse_workbook


@dataclass
class ExportError(Exception):
    """可映射到稳定 Worker 错误码的导出异常。"""

    code: str
    message: str


def calculate_file_fingerprint(file_path: str | Path) -> dict[str, Any]:
    """计算小型 Excel 源文件的稳定身份信息。"""

    path = Path(file_path).expanduser().resolve()
    if not path.is_file():
        raise ExportError("FILE_NOT_FOUND", "源文件不存在")

    digest = hashlib.sha256()
    try:
        with path.open("rb") as source:
            for chunk in iter(lambda: source.read(1024 * 1024), b""):
                digest.update(chunk)
        stat = path.stat()
    except PermissionError as error:
        raise ExportError("FILE_LOCKED", "源文件被占用或无权读取") from error

    return {
        "sha256": digest.hexdigest(),
        "sizeBytes": int(stat.st_size),
        "mtimeNs": int(stat.st_mtime_ns),
    }


def _require_task_id(payload: dict[str, Any]) -> str:
    task_id = payload.get("taskId")
    if not isinstance(task_id, str) or not task_id.strip():
        raise ExportError("INVALID_MESSAGE", "任务 ID 不能为空")
    return task_id.strip()


def _load_context(connection, task_id: str):
    task = connection.execute(
        "SELECT * FROM tasks WHERE id = ?",
        (task_id,),
    ).fetchone()
    if task is None:
        raise ExportError("TASK_NOT_FOUND", "翻译任务不存在")
    rows = connection.execute(
        """
        SELECT *
        FROM task_rows
        WHERE task_id = ?
        ORDER BY excel_row_number, row_id
        """,
        (task_id,),
    ).fetchall()
    return task, rows


def _validate_ready(task, rows) -> None:
    if task["status"] != "completed":
        raise ExportError(
            "EXPORT_TASK_INCOMPLETE",
            "仅已完成任务可以导出",
        )
    if any(row["status"] not in {"completed", "ignored"} for row in rows):
        raise ExportError(
            "EXPORT_TASK_INCOMPLETE",
            "仍有未确认、待填写或失败行",
        )


def _validate_source(task) -> tuple[dict[str, Any], dict[str, Any]]:
    expected_sha256 = task["source_file_sha256"]
    if not expected_sha256:
        raise ExportError(
            "SOURCE_FINGERPRINT_MISSING",
            "该任务创建于 M3 之前，请重新导入源文件后再导出",
        )

    current = calculate_file_fingerprint(task["source_file_path"])
    if (
        current["sha256"] != expected_sha256
        or current["sizeBytes"] != int(task["source_file_size_bytes"])
    ):
        raise ExportError(
            "SOURCE_FILE_CHANGED",
            "源文件自任务创建后已发生变化，请重新导入",
        )

    summary = parse_workbook(task["source_file_path"])
    if bool(task["source_restricted"]) or summary["restricted"]:
        raise ExportError(
            "EXPORT_RESTRICTED",
            "源文件包含高风险 Excel 对象，只允许只读预览",
        )
    return current, summary


def _resolve_target_column(task, summary: dict[str, Any]) -> tuple[int, bool]:
    sheet = next(
        (
            item
            for item in summary["sheets"]
            if item["name"] == task["sheet_name"]
        ),
        None,
    )
    if sheet is None:
        raise ExportError("SHEET_NOT_FOUND", "任务工作表不存在")

    existing = task["target_column"]
    if existing is not None:
        return int(existing), False

    target_column = int(sheet["columnCount"]) + 1
    if target_column > MAX_COLUMNS:
        raise ExportError(
            "COLUMN_LIMIT_EXCEEDED",
            "数据区右侧没有可用的新译文列",
        )
    return target_column, True


def _suggested_file_name(source_file_name: str) -> str:
    timestamp = datetime.now().astimezone().strftime("%Y%m%d_%H%M")
    return f"{Path(source_file_name).stem}_已翻译_{timestamp}.xlsx"


def _build_preflight(task, rows, summary: dict[str, Any]) -> dict[str, Any]:
    target_column, creates_target_column = _resolve_target_column(
        task,
        summary,
    )
    ignored_rows = sum(row["status"] == "ignored" for row in rows)
    unchanged_rows = sum(
        row["status"] == "completed"
        and row["translation"] is not None
        and json.loads(row["existing_target_json"])
        == row["translation"]
        for row in rows
    )
    writable_rows = sum(
        row["status"] == "completed"
        and isinstance(row["translation"], str)
        and bool(row["translation"].strip())
        for row in rows
    ) - unchanged_rows

    return {
        "taskId": task["id"],
        "ready": True,
        "sourceFilePath": task["source_file_path"],
        "sourceFileName": task["source_file_name"],
        "sourceSha256": task["source_file_sha256"],
        "sheetName": task["sheet_name"],
        "targetColumn": target_column,
        "targetColumnLetter": get_column_letter(target_column),
        "createsTargetColumn": creates_target_column,
        "writableRows": writable_rows,
        "unchangedRows": unchanged_rows,
        "ignoredRows": ignored_rows,
        "suggestedFileName": _suggested_file_name(
            task["source_file_name"],
        ),
        "risks": summary["risks"],
    }


def preflight_export(payload: dict[str, Any]) -> dict[str, Any]:
    """验证任务完成度、源文件身份和导出列边界。"""

    task_id = _require_task_id(payload)
    with connect_database() as connection:
        task, rows = _load_context(connection, task_id)
        _validate_ready(task, rows)
        _current, summary = _validate_source(task)
        return _build_preflight(task, rows, summary)


def _resolve_output_path(raw_path: Any, source_path: Path) -> Path:
    if not isinstance(raw_path, str) or not raw_path.strip():
        raise ExportError("EXPORT_PATH_INVALID", "输出路径不能为空")

    output_path = Path(raw_path).expanduser().resolve()
    if output_path.suffix.lower() != ".xlsx":
        raise ExportError("EXPORT_PATH_INVALID", "输出文件必须是 .xlsx")
    if os.path.normcase(str(output_path)) == os.path.normcase(str(source_path)):
        raise ExportError("EXPORT_PATH_INVALID", "输出文件不能覆盖原文件")
    if output_path.exists():
        raise ExportError("EXPORT_PATH_EXISTS", "输出文件已存在，请重新命名")
    if not output_path.parent.is_dir():
        raise ExportError("EXPORT_PATH_INVALID", "输出目录不存在")
    return output_path


def _copy_new_column_style(worksheet, source_column: int, target_column: int) -> None:
    source_letter = get_column_letter(source_column)
    target_letter = get_column_letter(target_column)
    source_width = worksheet.column_dimensions[source_letter].width
    worksheet.column_dimensions[target_letter].width = max(
        source_width or 13,
        20,
    )


def _copy_cell_style(source_cell, target_cell) -> None:
    if source_cell.has_style:
        target_cell._style = copy(source_cell._style)
    target_cell.number_format = source_cell.number_format
    target_cell.alignment = copy(source_cell.alignment)


def _write_rows(workbook, task, rows, preflight: dict[str, Any]):
    if task["sheet_name"] not in workbook.sheetnames:
        raise ExportError("SHEET_NOT_FOUND", "任务工作表不存在")
    worksheet = workbook[task["sheet_name"]]
    target_column = int(preflight["targetColumn"])
    creates_target_column = bool(preflight["createsTargetColumn"])
    source_column = int(task["source_column"])

    if creates_target_column:
        _copy_new_column_style(worksheet, source_column, target_column)
        header_source = worksheet.cell(
            row=int(task["header_row"]),
            column=source_column,
        )
        header_target = worksheet.cell(
            row=int(task["header_row"]),
            column=target_column,
        )
        _copy_cell_style(header_source, header_target)
        header_target.value = "中文翻译"

    written: list[tuple[str, str]] = []
    skipped_rows = 0
    for row in rows:
        if row["status"] == "ignored":
            skipped_rows += 1
            continue
        translation = row["translation"]
        if not isinstance(translation, str) or not translation.strip():
            raise ExportError(
                "EXPORT_TASK_INCOMPLETE",
                f"第 {row['excel_row_number']} 行译文为空",
            )
        existing_target = json.loads(row["existing_target_json"])
        if existing_target == translation:
            skipped_rows += 1
            continue

        target_cell = worksheet.cell(
            row=int(row["excel_row_number"]),
            column=target_column,
        )
        if isinstance(target_cell, MergedCell):
            raise ExportError(
                "EXPORT_CELL_MERGED",
                f"第 {row['excel_row_number']} 行目标单元格属于合并区域",
            )
        if creates_target_column:
            source_cell = worksheet.cell(
                row=int(row["excel_row_number"]),
                column=source_column,
            )
            _copy_cell_style(source_cell, target_cell)
        target_cell.value = translation
        written.append((target_cell.coordinate, translation))
    return written, skipped_rows


def _validate_output(temp_path: Path, sheet_name: str, written) -> None:
    try:
        workbook = load_workbook(
            temp_path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as error:
        raise ExportError(
            "EXPORT_VALIDATION_FAILED",
            "导出文件无法重新打开",
        ) from error

    try:
        if sheet_name not in workbook.sheetnames:
            raise ExportError(
                "EXPORT_VALIDATION_FAILED",
                "导出文件缺少任务工作表",
            )
        worksheet = workbook[sheet_name]
        for coordinate, translation in written:
            if worksheet[coordinate].value != translation:
                raise ExportError(
                    "EXPORT_VALIDATION_FAILED",
                    f"导出单元格 {coordinate} 校验失败",
                )
    finally:
        workbook.close()


def export_translation_task(payload: dict[str, Any]) -> dict[str, Any]:
    """写入临时副本，重新打开校验后再生成最终 Excel 文件。"""

    task_id = _require_task_id(payload)
    with connect_database() as connection:
        task, rows = _load_context(connection, task_id)
        _validate_ready(task, rows)
        _current, summary = _validate_source(task)
        preflight = _build_preflight(task, rows, summary)

    source_path = Path(task["source_file_path"]).resolve()
    output_path = _resolve_output_path(payload.get("outputPath"), source_path)
    temp_path = output_path.with_name(
        f".{output_path.stem}.{uuid4().hex}.tmp.xlsx"
    )

    workbook = None
    try:
        workbook = load_workbook(
            source_path,
            read_only=False,
            data_only=False,
            keep_links=False,
        )
        written, skipped_rows = _write_rows(
            workbook,
            task,
            rows,
            preflight,
        )
        workbook.save(temp_path)
        workbook.close()
        workbook = None
        _validate_output(temp_path, task["sheet_name"], written)
        temp_path.replace(output_path)
    except ExportError:
        raise
    except PermissionError as error:
        raise ExportError(
            "EXPORT_WRITE_FAILED",
            "输出文件或目录不可写",
        ) from error
    except OSError as error:
        raise ExportError("EXPORT_WRITE_FAILED", "写入输出文件失败") from error
    except Exception as error:
        raise ExportError(
            "EXPORT_WRITE_FAILED",
            "生成 Excel 副本失败",
        ) from error
    finally:
        if workbook is not None:
            workbook.close()
        if temp_path.exists():
            temp_path.unlink(missing_ok=True)

    output_fingerprint = calculate_file_fingerprint(output_path)
    return {
        "taskId": task_id,
        "outputPath": str(output_path),
        "outputFileName": output_path.name,
        "outputSha256": output_fingerprint["sha256"],
        "sheetName": task["sheet_name"],
        "targetColumn": preflight["targetColumn"],
        "targetColumnLetter": preflight["targetColumnLetter"],
        "createdTargetColumn": preflight["createsTargetColumn"],
        "writtenRows": len(written),
        "skippedRows": skipped_rows,
        "validated": True,
    }
