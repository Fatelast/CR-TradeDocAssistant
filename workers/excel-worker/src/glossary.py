"""M4 术语管理与标准 XLSX 导入导出。"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from database import build_glossary_term_key, connect_database
from exporter import calculate_file_fingerprint

MAX_GLOSSARY_FILE_BYTES = 5 * 1024 * 1024
MAX_GLOSSARY_ROWS = 5_000
TEMPLATE_HEADERS = (
    "俄文原文",
    "中文译文",
    "分类",
    "精确匹配",
    "区分大小写",
    "启用",
    "备注",
)


@dataclass
class GlossaryError(Exception):
    code: str
    message: str


def _now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def _get_version(connection: Any) -> int:
    row = connection.execute(
        "SELECT value FROM app_meta WHERE key = 'glossary_version'"
    ).fetchone()
    return int(row["value"]) if row else 1


def _increment_version(connection: Any) -> int:
    version = _get_version(connection) + 1
    connection.execute(
        """
        INSERT INTO app_meta(key, value) VALUES('glossary_version', ?)
        ON CONFLICT(key) DO UPDATE SET value = excluded.value
        """,
        (str(version),),
    )
    return version


def _require_text(value: Any, message: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise GlossaryError("INVALID_MESSAGE", message)
    return value.strip()


def _serialize_term(row: Any) -> dict[str, Any]:
    return {
        "termId": row["id"],
        "sourceText": row["source_text"],
        "targetText": row["target_text"],
        "category": row["category"],
        "exactMatch": bool(row["exact_match"]),
        "caseSensitive": bool(row["case_sensitive"]),
        "enabled": bool(row["enabled"]),
        "note": row["note"],
        "createdAt": row["created_at"],
        "updatedAt": row["updated_at"],
    }


def list_glossary_terms(payload: dict[str, Any]) -> dict[str, Any]:
    page = payload.get("page", 1)
    page_size = payload.get("pageSize", 20)
    search = payload.get("search", "")
    category = payload.get("category")
    enabled = payload.get("enabled")
    if (
        not isinstance(page, int)
        or isinstance(page, bool)
        or page < 1
        or not isinstance(page_size, int)
        or isinstance(page_size, bool)
        or page_size < 1
        or page_size > 50
        or not isinstance(search, str)
        or len(search) > 100
        or category is not None
        and (not isinstance(category, str) or not category.strip())
        or enabled not in {None, True, False}
    ):
        raise GlossaryError("INVALID_MESSAGE", "术语查询参数无效")

    clauses: list[str] = []
    parameters: list[Any] = []
    if search.strip():
        escaped = (
            search.strip()
            .replace("\\", "\\\\")
            .replace("%", "\\%")
            .replace("_", "\\_")
        )
        clauses.append(
            "(source_text LIKE ? ESCAPE '\\' "
            "OR target_text LIKE ? ESCAPE '\\')"
        )
        parameters.extend([f"%{escaped}%", f"%{escaped}%"])
    if category:
        clauses.append("category = ?")
        parameters.append(category.strip())
    if enabled is not None:
        clauses.append("enabled = ?")
        parameters.append(int(enabled))
    where_clause = (
        f"WHERE {' AND '.join(clauses)}"
        if clauses
        else ""
    )
    with connect_database() as connection:
        total = int(
            connection.execute(
                f"""
                SELECT COUNT(*) AS count FROM glossary_terms
                {where_clause}
                """,
                parameters,
            ).fetchone()["count"]
        )
        rows = connection.execute(
            f"""
            SELECT * FROM glossary_terms
            {where_clause}
            ORDER BY enabled DESC, category, source_text, id
            LIMIT ? OFFSET ?
            """,
            [*parameters, page_size, (page - 1) * page_size],
        ).fetchall()
        categories = [
            str(row["category"])
            for row in connection.execute(
                """
                SELECT DISTINCT category FROM glossary_terms
                ORDER BY category
                """
            ).fetchall()
        ]
        version = _get_version(connection)
    return {
        "glossaryVersion": version,
        "items": [_serialize_term(row) for row in rows],
        "categories": categories,
        "page": page,
        "pageSize": page_size,
        "total": total,
        "totalPages": max(1, (total + page_size - 1) // page_size),
    }


def upsert_glossary_term(payload: dict[str, Any]) -> dict[str, Any]:
    source_text = _require_text(payload.get("sourceText"), "术语原文不能为空")
    target_text = _require_text(payload.get("targetText"), "术语译文不能为空")
    category = str(payload.get("category") or "通用").strip() or "通用"
    exact_match = payload.get("exactMatch") is not False
    case_sensitive = payload.get("caseSensitive") is True
    enabled = payload.get("enabled") is not False
    note = payload.get("note")
    if note is not None and not isinstance(note, str):
        raise GlossaryError("INVALID_MESSAGE", "术语备注无效")
    if any(
        isinstance(value, str) and value.lstrip().startswith("=")
        for value in (source_text, target_text, category, note)
    ):
        raise GlossaryError(
            "INVALID_MESSAGE",
            "术语内容不能使用 Excel 公式",
        )
    term_id = payload.get("termId")
    if term_id is not None:
        term_id = _require_text(term_id, "术语 ID 无效")
    else:
        term_id = str(uuid4())
    term_key = build_glossary_term_key(
        source_text,
        exact_match,
        case_sensitive,
    )
    now = _now()

    with connect_database() as connection:
        duplicate = connection.execute(
            """
            SELECT id FROM glossary_terms
            WHERE term_key = ? AND id <> ?
            """,
            (term_key, term_id),
        ).fetchone()
        if duplicate is not None:
            raise GlossaryError(
                "GLOSSARY_TERM_DUPLICATE",
                "相同匹配规则的术语已经存在",
            )
        existing = connection.execute(
            "SELECT id FROM glossary_terms WHERE id = ?",
            (term_id,),
        ).fetchone()
        if existing is None:
            connection.execute(
                """
                INSERT INTO glossary_terms(
                    id, source_text, target_text, category,
                    exact_match, case_sensitive, enabled, note,
                    created_at, updated_at, term_key
                ) VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    term_id,
                    source_text,
                    target_text,
                    category,
                    int(exact_match),
                    int(case_sensitive),
                    int(enabled),
                    note,
                    now,
                    now,
                    term_key,
                ),
            )
        else:
            connection.execute(
                """
                UPDATE glossary_terms
                SET source_text = ?, target_text = ?, category = ?,
                    exact_match = ?, case_sensitive = ?, enabled = ?,
                    note = ?, updated_at = ?, term_key = ?
                WHERE id = ?
                """,
                (
                    source_text,
                    target_text,
                    category,
                    int(exact_match),
                    int(case_sensitive),
                    int(enabled),
                    note,
                    now,
                    term_key,
                    term_id,
                ),
            )
        version = _increment_version(connection)
    return {"termId": term_id, "glossaryVersion": version}


def _parse_bool(value: Any, row_number: int, label: str) -> bool:
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().casefold()
    if normalized in {"是", "true", "1", "启用"}:
        return True
    if normalized in {"否", "false", "0", "停用"}:
        return False
    raise GlossaryError(
        "GLOSSARY_IMPORT_INVALID",
        f"第 {row_number} 行的“{label}”必须填写是或否",
    )


def _read_import_rows(file_path: Path) -> list[dict[str, Any]]:
    if not file_path.is_file():
        raise GlossaryError("FILE_NOT_FOUND", "术语文件不存在")
    if file_path.suffix.lower() != ".xlsx":
        raise GlossaryError("FILE_UNSUPPORTED", "术语只支持 XLSX 文件")
    if file_path.stat().st_size > MAX_GLOSSARY_FILE_BYTES:
        raise GlossaryError(
            "GLOSSARY_IMPORT_LIMIT_EXCEEDED",
            "术语文件不能超过 5 MiB",
        )
    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=False,
        keep_links=False,
    )
    try:
        worksheet = workbook.active
        headers = tuple(
            worksheet.cell(1, index).value
            for index in range(1, len(TEMPLATE_HEADERS) + 1)
        )
        if headers != TEMPLATE_HEADERS:
            raise GlossaryError(
                "GLOSSARY_IMPORT_INVALID",
                "术语文件表头与标准模板不一致",
            )
        if worksheet.max_row - 1 > MAX_GLOSSARY_ROWS:
            raise GlossaryError(
                "GLOSSARY_IMPORT_LIMIT_EXCEEDED",
                "术语数量不能超过 5,000 条",
            )
        items: list[dict[str, Any]] = []
        seen_keys: set[str] = set()
        for row_number in range(2, worksheet.max_row + 1):
            values = [
                worksheet.cell(row_number, column).value
                for column in range(1, len(TEMPLATE_HEADERS) + 1)
            ]
            if all(value is None or str(value).strip() == "" for value in values):
                continue
            source_text = str(values[0] or "").strip()
            target_text = str(values[1] or "").strip()
            contains_formula = any(
                isinstance(value, str)
                and value.lstrip().startswith("=")
                for value in values
            )
            if not source_text or not target_text or contains_formula:
                raise GlossaryError(
                    "GLOSSARY_IMPORT_INVALID",
                    f"第 {row_number} 行的俄文原文或中文译文无效",
                )
            exact_match = _parse_bool(values[3], row_number, "精确匹配")
            case_sensitive = _parse_bool(
                values[4],
                row_number,
                "区分大小写",
            )
            enabled = _parse_bool(values[5], row_number, "启用")
            term_key = build_glossary_term_key(
                source_text,
                exact_match,
                case_sensitive,
            )
            if term_key in seen_keys:
                raise GlossaryError(
                    "GLOSSARY_TERM_DUPLICATE",
                    f"第 {row_number} 行存在重复术语",
                )
            seen_keys.add(term_key)
            items.append(
                {
                    "termKey": term_key,
                    "sourceText": source_text,
                    "targetText": target_text,
                    "category": str(values[2] or "通用").strip() or "通用",
                    "exactMatch": exact_match,
                    "caseSensitive": case_sensitive,
                    "enabled": enabled,
                    "note": (
                        str(values[6]).strip()
                        if values[6] is not None
                        else None
                    ),
                }
            )
        return items
    finally:
        workbook.close()


def _build_import_plan(file_path: Path) -> dict[str, Any]:
    items = _read_import_rows(file_path)
    with connect_database() as connection:
        existing = {
            row["term_key"]: row
            for row in connection.execute(
                "SELECT * FROM glossary_terms WHERE term_key IS NOT NULL"
            ).fetchall()
        }
    created = 0
    updated = 0
    unchanged = 0
    for item in items:
        current = existing.get(item["termKey"])
        if current is None:
            created += 1
            continue
        comparable = (
            current["source_text"],
            current["target_text"],
            current["category"],
            bool(current["exact_match"]),
            bool(current["case_sensitive"]),
            bool(current["enabled"]),
            current["note"],
        )
        incoming = (
            item["sourceText"],
            item["targetText"],
            item["category"],
            item["exactMatch"],
            item["caseSensitive"],
            item["enabled"],
            item["note"],
        )
        if comparable == incoming:
            unchanged += 1
        else:
            updated += 1
    fingerprint = calculate_file_fingerprint(file_path)
    return {
        "fileName": file_path.name,
        "fileSha256": fingerprint["sha256"],
        "totalRows": len(items),
        "createdRows": created,
        "updatedRows": updated,
        "unchangedRows": unchanged,
        "items": items,
    }


def preflight_glossary_import(payload: dict[str, Any]) -> dict[str, Any]:
    file_path = payload.get("filePath")
    if not isinstance(file_path, str) or not file_path.strip():
        raise GlossaryError("INVALID_MESSAGE", "术语文件路径不能为空")
    plan = _build_import_plan(Path(file_path))
    return {key: value for key, value in plan.items() if key != "items"}


def apply_glossary_import(payload: dict[str, Any]) -> dict[str, Any]:
    file_path = payload.get("filePath")
    expected_sha = payload.get("expectedSha256")
    if (
        not isinstance(file_path, str)
        or not file_path.strip()
        or not isinstance(expected_sha, str)
        or not expected_sha
    ):
        raise GlossaryError("INVALID_MESSAGE", "术语导入请求无效")
    plan = _build_import_plan(Path(file_path))
    if plan["fileSha256"] != expected_sha:
        raise GlossaryError(
            "GLOSSARY_IMPORT_CHANGED",
            "术语文件在预检后发生变化",
        )
    now = _now()
    changed = plan["createdRows"] + plan["updatedRows"]
    with connect_database() as connection:
        existing = {
            row["term_key"]: row["id"]
            for row in connection.execute(
                "SELECT id, term_key FROM glossary_terms"
            ).fetchall()
        }
        for item in plan["items"]:
            term_id = existing.get(item["termKey"])
            if term_id is None:
                connection.execute(
                    """
                    INSERT INTO glossary_terms(
                        id, source_text, target_text, category,
                        exact_match, case_sensitive, enabled, note,
                        created_at, updated_at, term_key
                    ) VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        str(uuid4()),
                        item["sourceText"],
                        item["targetText"],
                        item["category"],
                        int(item["exactMatch"]),
                        int(item["caseSensitive"]),
                        int(item["enabled"]),
                        item["note"],
                        now,
                        now,
                        item["termKey"],
                    ),
                )
            else:
                connection.execute(
                    """
                    UPDATE glossary_terms
                    SET source_text = ?, target_text = ?, category = ?,
                        exact_match = ?, case_sensitive = ?, enabled = ?,
                        note = ?, updated_at = ?
                    WHERE id = ?
                    """,
                    (
                        item["sourceText"],
                        item["targetText"],
                        item["category"],
                        int(item["exactMatch"]),
                        int(item["caseSensitive"]),
                        int(item["enabled"]),
                        item["note"],
                        now,
                        term_id,
                    ),
                )
        version = _increment_version(connection) if changed else _get_version(
            connection
        )
    return {
        "glossaryVersion": version,
        "createdRows": plan["createdRows"],
        "updatedRows": plan["updatedRows"],
        "unchangedRows": plan["unchangedRows"],
    }


def export_glossary(payload: dict[str, Any]) -> dict[str, Any]:
    output_value = payload.get("outputPath")
    if not isinstance(output_value, str) or not output_value.strip():
        raise GlossaryError("EXPORT_PATH_INVALID", "术语导出路径不能为空")
    output_path = Path(output_value)
    if not output_path.is_absolute() or output_path.suffix.lower() != ".xlsx":
        raise GlossaryError("EXPORT_PATH_INVALID", "术语只能导出为 XLSX")
    if output_path.exists():
        raise GlossaryError("EXPORT_PATH_EXISTS", "术语导出文件已存在")
    if not output_path.parent.is_dir():
        raise GlossaryError("EXPORT_PATH_INVALID", "术语导出目录不存在")

    with connect_database() as connection:
        rows = connection.execute(
            """
            SELECT * FROM glossary_terms
            ORDER BY category, source_text, id
            """
        ).fetchall()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "术语库"
    header_fill = PatternFill("solid", fgColor="123C35")
    header_font = Font(color="F5EBD6", bold=True)
    for column, header in enumerate(TEMPLATE_HEADERS, 1):
        cell = worksheet.cell(1, column, header)
        cell.fill = copy(header_fill)
        cell.font = copy(header_font)
        cell.alignment = Alignment(horizontal="center")
    for row_number, row in enumerate(rows, 2):
        values = (
            row["source_text"],
            row["target_text"],
            row["category"],
            "是" if row["exact_match"] else "否",
            "是" if row["case_sensitive"] else "否",
            "是" if row["enabled"] else "否",
            row["note"],
        )
        for column, value in enumerate(values, 1):
            worksheet.cell(row_number, column, value)
    for column, width in zip("ABCDEFG", (32, 28, 14, 12, 14, 10, 32)):
        worksheet.column_dimensions[column].width = width
    worksheet.freeze_panes = "A2"
    temp_path = output_path.with_name(
        f".{output_path.stem}.{uuid4().hex}.tmp.xlsx"
    )
    try:
        workbook.save(temp_path)
        workbook.close()
        validation = load_workbook(temp_path, read_only=True, data_only=False)
        try:
            if tuple(
                validation.active.cell(1, index).value
                for index in range(1, len(TEMPLATE_HEADERS) + 1)
            ) != TEMPLATE_HEADERS:
                raise GlossaryError(
                    "EXPORT_VALIDATION_FAILED",
                    "术语导出文件校验失败",
                )
        finally:
            validation.close()
        temp_path.replace(output_path)
    except GlossaryError:
        raise
    except OSError as error:
        raise GlossaryError("EXPORT_WRITE_FAILED", "术语导出失败") from error
    finally:
        temp_path.unlink(missing_ok=True)

    return {
        "outputPath": str(output_path),
        "outputFileName": output_path.name,
        "termCount": len(rows),
        "validated": True,
    }
